from contextlib import asynccontextmanager
import io
import logging
from datetime import datetime

import pandas as pd
import re
import asyncpg
from fastapi import FastAPI, HTTPException, Depends, status, Request
from fastapi.exceptions import RequestValidationError
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import config
import database
import ml_engine
import schemas
import repositories
from safari_response import SafariFriendlyJSONResponse

TARGET_WEIGHT_ZONES = ["Овощи и фрукты", "Свежее мясо", "Мясной прилавок", "Мясной цех"]
BASE_TRANSIT_TIME_SECONDS = 15
SETUP_TIME_SECONDS = 240
MAX_EXCEL_TITLE_LEN = 31


@asynccontextmanager
async def lifespan(_: FastAPI):
    await database.init_db()
    async for db in database.get_db():
        await ml_engine.train_all_models_async(db)
        break
    yield


app = FastAPI(
    title="Warehouse Automation System",
    version="1.0.0",
    default_response_class=SafariFriendlyJSONResponse,
    lifespan=lifespan
)

logger = logging.getLogger(__name__)


@app.exception_handler(RequestValidationError)
async def validation_exception_handler(request: Request, exc: RequestValidationError):
    return SafariFriendlyJSONResponse(
        status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
        content={"status": "error", "message": "Ошибка валидации входящих данных", "details": exc.errors()}
    )


@app.exception_handler(HTTPException)
async def fastapi_http_exception_handler(request: Request, exc: HTTPException):
    return SafariFriendlyJSONResponse(
        status_code=exc.status_code,
        content={"status": "error", "message": exc.detail}
    )


@app.exception_handler(Exception)
async def global_exception_handler(request: Request, exc: Exception):
    logger.error(f"Ошибка БЭКЕНДА при обработке запроса [{request.method}] {request.url.path}: {exc}", exc_info=True)

    details = str(exc)
    if hasattr(exc, "errors"):
        try:
            details = exc.errors()
        except Exception:
            pass

    return SafariFriendlyJSONResponse(
        status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
        content={
            "status": "error",
            "message": "Внутренняя ошибка сервера при работе с логистикой/БД",
            "details": details,
            "developer": "Nadir"
        }
    )


def _get_current_time_metrics(stats_based_peak_days: list[int]) -> tuple[int, float, bool]:
    now = datetime.now()
    current_hour = now.hour
    human_weekday = now.weekday() + 1
    is_peak_day = human_weekday in stats_based_peak_days

    time_modifier = 1.0
    if is_peak_day:
        time_modifier = 1.5 if 16 <= current_hour <= 22 else 1.2
    elif 17 <= current_hour <= 20:
        time_modifier = 1.3
    return human_weekday, time_modifier, is_peak_day


def _calculate_zone_congestion(active_orders: list) -> dict:
    zone_congestion = {}
    for row in active_orders:
        item_zone = config.find_zone_for_item(row["item_name"])
        zone_congestion[item_zone] = zone_congestion.get(item_zone, 0) + row["active_count"]
    return zone_congestion


def _calculate_worker_verdict(cancel_rate: float, total_processed: int) -> tuple[str, str]:
    if cancel_rate > 25 and total_processed >= 5:
        return "Высокий процент отмен (требуется проверка полок)", "FCE4D6"
    if total_processed > 30:
        return "Лидер смены (высокая нагрузка)", "E2EFDA"
    return "Стабильная работа", "FFFFFF"


def _get_zone_base_modifier(zone: str) -> float:
    return config.ZONE_SPEED_MODIFIERS.get(zone, config.DEFAULT_PICK_TIME)


def _resolve_store_topology(store_id: str) -> dict:
    store_id_normalized = store_id.lower()
    route_graph = config.STORES_TOPOLOGIES.get(store_id_normalized)
    if not route_graph:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Store '{store_id}' has no registered topology"
        )
    return route_graph


@app.get('/')
def root():
    return {'project': 'hypermarket-logistics-ai', 'status': 'ready'}


@app.get('/zone/{zone_id}')
def check_zone(zone_id: int):
    return {'current_zone': zone_id, 'status': 'Зона доступна для сборки товаров'}


@app.get('/analytics/top-workers')
async def top_workers(db: asyncpg.Connection = Depends(database.get_db)):
    row = await repositories.AnalyticsRepository.get_top_worker(db)
    if not row:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Сотрудники не найдены в базе данных.")
    return {
        "status": "Аналитика успешно рассчитана через SQL-репозиторий",
        "лучший_сборщик": row["name"],
        "максимум_заказов": row["orders_per_shift"]
    }


@app.get('/analytics/canceled-stats')
async def get_canceled_stats(worker_name: str = None, db: asyncpg.Connection = Depends(database.get_db)):
    if worker_name:
        row = await repositories.AnalyticsRepository.get_canceled_stats_by_worker(db, worker_name)
        if row is None or row["worker_id"] is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND,
                                detail=f"Сборщик с именем '{worker_name}' не найден.")
        return {
            "status": f"Аналитика по сотруднику {worker_name} успешна",
            "сборщик": worker_name,
            "количество_отмен": row["cnt"],
            "личный_ущерб_рублей": row["total"] or 0
        }

    rows = await repositories.AnalyticsRepository.get_all_workers_canceled_counts(db)
    stats_dict = {row["name"]: row["cnt"] for row in rows}
    total_lost_money = await repositories.AnalyticsRepository.get_total_lost_money(db)

    return {
        "status": "Общая аналитика отмен успешна",
        "отмены_по_сборщикам": stats_dict,
        "общий_ущерб_рублей": total_lost_money or 0
    }


@app.post('/orders/optimize')
def optimize_order_route(products: list[schemas.ProductItem]):
    if not products:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Список товаров пуст")
    zones_map = {}
    for product in products:
        zones_map.setdefault(product.zone, []).append(product.name)
    return {"status": "Маршрут сборки успешно оптимизирован", "route": zones_map}


@app.post('/api/v1/orders/build-zone-route')
def build_zone_route(order: schemas.Supermarket):
    route_graph = _resolve_store_topology(order.store_id)
    processed_items = []
    for item in order.items:
        zone = config.find_zone_for_item(item)
        step_number = config.get_zone_step_for_store(route_graph, zone)
        processed_items.append({"item_name": item, "zone": zone, "route_step": step_number})
    return {
        "status": "success",
        "active_store": order.store_id.lower(),
        "total_items_in_order": len(order.items),
        "optimized_picking_route": sorted(processed_items, key=lambda x: (x["route_step"], x["zone"]))
    }


@app.post('/api/v1/orders/build-optimized-route')
async def build_optimized_route(order: schemas.Supermarket, db: asyncpg.Connection = Depends(database.get_db)):
    store_id = order.store_id.lower()
    route_graph = _resolve_store_topology(store_id)

    top_days_rows = await repositories.OrderRepository.get_peak_days_statistics(db, store_id)
    stats_based_peak_days = [row["day_of_week"] for row in top_days_rows]

    human_weekday, time_modifier, is_peak_day_by_stats = _get_current_time_metrics(stats_based_peak_days)

    active_orders = await repositories.OrderRepository.get_active_workers_congestion(db, store_id)
    zone_congestion = _calculate_zone_congestion(active_orders)

    zone_to_items = {}
    for item in order.items:
        zone = config.find_zone_for_item(item)
        zone_to_items.setdefault(zone, []).append(item)

    route_steps = []
    total_picking_time_seconds = 0
    for zone, items in zone_to_items.items():
        step_number = config.get_zone_step_for_store(route_graph, zone)
        base_modifier = _get_zone_base_modifier(zone)
        active_workers = zone_congestion.get(zone, 0)
        congestion_modifier = 1.0 if active_workers <= 2 else 1.3 if active_workers <= 5 else 1.6

        zone_weight = sum(config.extract_weight_kg(item) for item in items) if zone in TARGET_WEIGHT_ZONES else 0.0
        weight_penalty = max(0.0, zone_weight * 15.0)

        zone_picking_time = (len(items) * base_modifier * congestion_modifier * time_modifier) + weight_penalty
        total_picking_time_seconds += zone_picking_time
        route_steps.append({
            "zone": zone,
            "route_step": step_number,
            "items_count": len(items),
            "total_zone_weight_kg": round(zone_weight, 2),
            "weight_penalty_seconds": round(weight_penalty, 1),
            "items": items,
            "active_workers_here": active_workers,
            "congestion_coefficient": congestion_modifier,
            "estimated_zone_time_seconds": round(zone_picking_time, 1)
        })

    route_steps.sort(key=lambda x: (x["route_step"], x["zone"]))
    transit_time = (max(0, len(route_steps) - 1) * BASE_TRANSIT_TIME_SECONDS) + (
        SETUP_TIME_SECONDS if len(route_steps) > 0 else 0)
    total_time_seconds = total_picking_time_seconds + transit_time

    return {
        "status": "success",
        "store_id": store_id,
        "total_unique_zones_to_visit": len(route_steps),
        "total_items": len(order.items),
        "current_human_weekday": human_weekday,
        "detected_peak_days_by_statistics": stats_based_peak_days,
        "is_today_stat_peak": is_peak_day_by_stats,
        "external_hour_kick_modifier": round(time_modifier, 2),
        "estimated_total_picking_time_minutes": round(total_time_seconds / 60, 1),
        "optimized_route_steps": route_steps
    }


@app.post('/api/v1/orders/assign')
async def assign_order(order: schemas.NewOrder, db: asyncpg.Connection = Depends(database.get_db)):
    _resolve_store_topology(order.store_id)

    worker = await repositories.OrderRepository.find_least_loaded_worker(db)
    if not worker:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Нет доступных сборщиков.")

    target_worker_id, target_worker_name = worker["worker_id"], worker["name"]
    max_id = await repositories.OrderRepository.get_max_order_id(db)
    next_order_id = (max_id or 100000) + 1

    await repositories.OrderRepository.create_and_assign_order(
        db, next_order_id, order.store_id.lower(), order.item_name, order.item_price, target_worker_id
    )

    return {
        "status": "success",
        "message": "Заказ успешно создан и назначен через паттерн Repository",
        "order_id": next_order_id,
        "assigned_worker": {
            "worker_id": target_worker_id,
            "name": target_worker_name,
            "new_orders_count": worker["orders_per_shift"] + 1
        }
    }


@app.post('/api/v1/orders/suggest-substitution')
def suggest_substitution(request: schemas.SubstitutionRequest):
    detected_zone = config.find_zone_for_item(request.missing_item)
    alternatives = []
    rating = []

    if detected_zone != "Без зоны":
        raw_alternatives = [
            cat_item for cat_item, zone in config.ITEM_TO_ZONE.items()
            if zone == detected_zone and cat_item.lower() != request.missing_item.lower()
        ]

        if raw_alternatives:
            now = datetime.now()
            cur_day, cur_hour = now.weekday() + 1, now.hour
            scored_items = []

            for item in raw_alternatives:
                model = ml_engine.ML_MODELS_CACHE["availability_models"].get(item)
                if model:
                    x_train = pd.DataFrame([[cur_day, cur_hour]], columns=['day_of_week', 'hour'])
                    prob_canceled = model.predict_proba(x_train)
                    scored_items.append((item, float(prob_canceled[0][1])))
                else:
                    scored_items.append((item, 0.0))

            scored_items.sort(key=lambda x: x[1])
            alternatives = [item[0] for item in scored_items]
            rating = [
                {
                    "product_name": name,
                    "shelf_absence_probability": f"{round(prob * 100, 1)}%",
                    "confidence_status": "Высокая вероятность наличия" if prob < 0.3 else "Риск отсутствия на полке"
                } for name, prob in scored_items[:3]
            ]

    if not alternatives:
        return {
            "status": "success",
            "requested_item": request.missing_item,
            "detected_zone": detected_zone,
            "total_substitutions_found": 0,
            "predictive_availability_rating": [],
            "recommended_substitutions": [],
            "info": f"Альтернативные товары в зоне '{detected_zone}' не найдены."
        }

    return {
        "status": "success",
        "requested_item": request.missing_item,
        "detected_zone": detected_zone,
        "total_substitutions_found": len(alternatives),
        "predictive_availability_rating": rating,
        "recommended_substitutions": alternatives[:3]
    }


@app.get('/api/v1/analytics/predict-load')
def predict_store_load(store_id: str):
    if not store_id:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Идентификатор магазина обязателен.")

    _resolve_store_topology(store_id)

    model = ml_engine.ML_MODELS_CACHE["load_models"].get(store_id.lower())
    now = datetime.now()
    current_hour = now.hour
    human_weekday = now.weekday() + 1

    if model:
        x_train = pd.DataFrame([[human_weekday, current_hour]], columns=['day_of_week', 'hour'])
        predicted_orders = model.predict(x_train)[0]
        load_factor = max(1.0, round(float(predicted_orders) / 15.0, 2))
        reason = "Прогноз извлечен из RAM-кэша моделей LinearRegression (Production Pattern)"
    else:
        load_factor = 1.0
        reason = "Модель для данного магазина отсутствует в кэше, задействован базовый KPI"

    return {
        "status": "success",
        "store_id": store_id.lower(),
        "checked_at_hour": current_hour,
        "current_human_weekday": human_weekday,
        "kpi_time_modifier": load_factor,
        "verdict": f"Рекомендовано увеличить норматив времени сборки в {load_factor}x раз.",
        "reason": reason
    }


@app.get('/api/v1/admin/workers-efficiency')
async def get_workers_efficiency(store_id: str = None, db: asyncpg.Connection = Depends(database.get_db)):
    rows = await repositories.AnalyticsRepository.get_aggregated_efficiency_report(db, store_id)
    performance_report = []

    for row in rows:
        w_id = row["worker_id"]
        w_name = row["name"]
        success_cnt = row["success_cnt"]
        canceled_cnt = row["canceled_cnt"]

        total_processed = success_cnt + canceled_cnt
        cancel_rate = round((canceled_cnt / total_processed * 100), 1) if total_processed > 0 else 0.0

        verdict, _ = _calculate_worker_verdict(cancel_rate, total_processed)

        performance_report.append({
            "id": w_id, "name": w_name, "total_orders_processed": total_processed,
            "successful_orders": success_cnt, "canceled_orders": canceled_cnt,
            "cancel_rate": f"{cancel_rate}%", "system_verdict": verdict
        })

    return {
        "status": "success", "filter_by_store": store_id or "Все магазины",
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S Local"),
        "total_workers_tracked": len(performance_report), "report": performance_report
    }


@app.get('/api/v1/admin/workers-efficiency/excel')
async def get_workers_efficiency_excel(store_id: str = None, db: asyncpg.Connection = Depends(database.get_db)):
    rows = await repositories.AnalyticsRepository.get_aggregated_efficiency_report(db, store_id)

    wb = Workbook()
    ws = wb.active

    raw_title = f"Смена - {store_id or 'Общая'}"
    ws.title = re.sub(r'[\\*?:/\[\]]', '', raw_title)[:MAX_EXCEL_TITLE_LEN]
    ws.sheet_view.showGridLines = True

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    regular_font = Font(name="Calibri", size=11)
    bold_font = Font(name="Calibri", size=11, bold=True)

    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    thin_side = Side(border_style="thin", color="D9D9D9")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    headers = ["ID сборщика", "Имя сотрудника", "Всего заказов", "Успешно собрано", "Отменено позиций",
               "Процент отмен", "Вердикт системы"]
    ws.append(headers)

    for col_num, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    for row in rows:
        w_id = row["worker_id"]
        w_name = row["name"]
        success_cnt = row["success_cnt"]
        canceled_cnt = row["canceled_cnt"]

        total_processed = success_cnt + canceled_cnt
        cancel_rate = round((canceled_cnt / total_processed * 100), 1) if total_processed > 0 else 0.0

        verdict, verdict_color = _calculate_worker_verdict(cancel_rate, total_processed)

        row_idx = ws.max_row + 1
        ws.append([w_id, w_name, total_processed, success_cnt, canceled_cnt, f"{cancel_rate}%", verdict])

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_num)
            cell.font = regular_font
            cell.border = thin_border
            val_str = str(cell.value or "")

            if val_str.isdigit() or "%" in val_str:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

            if col_num == len(headers) and verdict_color != "FFFFFF":
                cell.fill = PatternFill(start_color=verdict_color, end_color=verdict_color, fill_type="solid")
                cell.font = bold_font

    ws.row_dimensions.height = 26
    for row_num in range(2, ws.max_row + 1):
        ws.row_dimensions[row_num].height = 20

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)

    filename_store = f"efficiency_report_{store_id.lower() if store_id else 'all'}.xlsx"
    return StreamingResponse(
        file_stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename_store}"}
    )


@app.post('/api/v1/orders/build-multi-optimized-route')
async def build_multi_optimized_route(request: schemas.MultiOrderRequest,
                                      db: asyncpg.Connection = Depends(database.get_db)):
    store_id = request.store_id.lower()
    route_graph = _resolve_store_topology(store_id)

    top_days_rows = await repositories.OrderRepository.get_peak_days_statistics(db, store_id)
    stats_based_peak_days = [row["day_of_week"] for row in top_days_rows]

    human_weekday, time_modifier, is_peak_day_by_stats = _get_current_time_metrics(stats_based_peak_days)

    active_orders = await repositories.OrderRepository.get_active_workers_congestion(db, store_id)
    zone_congestion = _calculate_zone_congestion(active_orders)

    zone_to_manifest = {}
    total_items_count = 0

    for order_idx, basket in enumerate(request.orders, 1):
        for item in basket:
            total_items_count += 1
            zone = config.find_zone_for_item(item)
            zone_to_manifest.setdefault(zone, []).append({"client_id": f"Заказ #{order_idx}", "item_name": item})

    route_steps = []
    total_picking_time_seconds = 0

    for zone, manifests in zone_to_manifest.items():
        step_number = config.get_zone_step_for_store(route_graph, zone)
        base_modifier = config.ZONE_SPEED_MODIFIERS.get(zone, config.DEFAULT_PICK_TIME)
        active_workers = zone_congestion.get(zone, 0)
        congestion_modifier = 1.0 if active_workers <= 2 else 1.3 if active_workers <= 5 else 1.6

        zone_weight = sum(config.extract_weight_kg(m["item_name"]) for m in manifests) \
            if zone in TARGET_WEIGHT_ZONES else 0.0
        weight_penalty = max(0.0, zone_weight * 15.0)

        zone_picking_time = (len(manifests) * base_modifier * congestion_modifier * time_modifier) + weight_penalty
        total_picking_time_seconds += zone_picking_time

        route_steps.append({
            "zone": zone, "route_step": step_number, "total_items_here_count": len(manifests),
            "total_zone_weight_kg": round(zone_weight, 2), "weight_penalty_seconds": round(weight_penalty, 1),
            "picking_manifest": manifests, "active_workers_here": active_workers,
            "congestion_coefficient": congestion_modifier, "estimated_zone_time_seconds": round(zone_picking_time, 1)
        })

    route_steps.sort(key=lambda x: (x["route_step"], x["zone"]))
    transit_time = (max(0, len(route_steps) - 1) * BASE_TRANSIT_TIME_SECONDS) + \
                   (SETUP_TIME_SECONDS if len(route_steps) > 0 else 0)
    total_time_seconds = total_picking_time_seconds + transit_time

    return {
        "status": "success", "store_id": store_id, "total_combined_orders": len(request.orders),
        "total_items_to_pick": total_items_count, "current_human_weekday": human_weekday,
        "detected_peak_days_by_statistics": stats_based_peak_days, "is_today_stat_peak": is_peak_day_by_stats,
        "external_hour_kick_modifier": round(time_modifier, 2),
        "estimated_total_multi_picking_time_minutes": round(total_time_seconds / 60, 1),
        "optimized_combined_route_steps": route_steps
    }